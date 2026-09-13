import csv
import hashlib
import json
import os
import shutil
from pathlib import Path, PureWindowsPath

from docx import Document
from openpyxl import load_workbook
from pptx import Presentation
from pypdf import PdfReader
from send2trash import send2trash

from config.app_settings import load_settings
from security.permissions import (
    is_admin,
    is_sensitive_path,
    validate_delete,
    validate_read,
    validate_write,
)


TEXT_EXTENSIONS = {".txt", ".md", ".py", ".log", ".csv", ".json"}
SKIP_SEARCH_DIRECTORIES = {".venv", ".git", "__pycache__", "node_modules"}
PREVIEW_LIMIT = 3000


def resolve_user_path(path: str) -> Path:
    home = Path.home()
    path = str(path).strip().strip('"').strip("'")
    raw_path = Path(path)

    if raw_path.is_absolute():
        return raw_path.expanduser().resolve()

    parts = list(PureWindowsPath(path).parts)
    known_folders = {
        "desktop": "Desktop",
        "masaüstü": "Desktop",
        "masaustu": "Desktop",
        "documents": "Documents",
        "belgeler": "Documents",
        "downloads": "Downloads",
        "indirilenler": "Downloads",
    }

    for index, part in enumerate(parts):
        folder = known_folders.get(part.lower())
        if folder:
            resolved = home / folder
            for remaining in parts[index + 1:]:
                resolved /= remaining
            return resolved.resolve()

    candidates = [
        home / "Desktop" / raw_path,
        home / "Documents" / raw_path,
        home / "Downloads" / raw_path,
    ]

    for candidate in candidates:
        if candidate.exists():
            return candidate.resolve()

    return (home / "Desktop" / raw_path).resolve()


def _file_fingerprint(path: Path) -> str:
    sha256 = hashlib.sha256()

    with path.open("rb") as file:
        while chunk := file.read(65536):
            sha256.update(chunk)

    return sha256.hexdigest()


def _truncate_text(text: str, limit: int = PREVIEW_LIMIT) -> str:
    if len(text) <= limit:
        return text

    return text[:limit] + "\n...[preview shortened]"


def _can_use_open_path(source: str, user_id: int | None) -> bool:
    return source == "terminal" or (
        source == "web"
        and is_admin(source, user_id)
    )


def _can_send_telegram_file(source: str, user_id: int | None) -> bool:
    return source == "terminal" or (
        source in {"web", "telegram"}
        and is_admin(source, user_id)
    )


def _iter_search_files(root: Path):
    if root.is_file():
        yield root
        return

    for current_root, directories, files in os.walk(root):
        directories[:] = [
            name
            for name in directories
            if name.lower() not in SKIP_SEARCH_DIRECTORIES
        ]

        current_path = Path(current_root)

        for filename in files:
            yield current_path / filename


def _search_roots(source: str, user_id: int | None):
    home = Path.home()
    roots = [
        home / "Desktop",
        home / "Documents",
        home / "Downloads",
    ]

    if source != "terminal":
        return roots

    try:
        local_admin = load_settings().get("local_admin", {})
    except Exception:
        return roots

    if (
        local_admin.get("access_mode") != "custom"
        or not local_admin.get("custom_access_acknowledged")
    ):
        return roots

    for value in local_admin.get("allowed_paths", []):
        try:
            root = validate_read(
                resolve_user_path(value),
                source,
                user_id,
            )
        except (OSError, ValueError, PermissionError):
            continue

        if root.is_dir() and root not in roots:
            roots.append(root)

    return roots


def _confirmation(
    action,
    path,
    title,
    *,
    destination=None,
    mode=None,
    content=None,
    old_preview=None,
    new_preview=None,
):
    result = {
        "status": "confirmation_required",
        "action": action,
        "path": str(path),
        "fingerprint": _file_fingerprint(path),
        "title": title,
    }

    if destination is not None:
        result["destination"] = str(destination)

    if mode is not None:
        result["mode"] = mode

    if content is not None:
        result["content"] = content

    if old_preview is not None:
        result["old_preview"] = old_preview

    if new_preview is not None:
        result["new_preview"] = new_preview

    return result


def create_file(
    path: str,
    content: str = "",
    _source: str = "terminal",
    _user_id: int | None = None,
):
    file_path = validate_write(
        resolve_user_path(path),
        _source,
        _user_id,
    )

    if file_path.exists():
        if not file_path.is_file():
            raise ValueError("The provided path is not a file.")

        if file_path.suffix.lower() not in TEXT_EXTENSIONS:
            raise ValueError(
                "Text cannot be written over this existing file type."
            )

        old_content = file_path.read_text(
            encoding="utf-8",
            errors="replace",
        )

        if old_content == content:
            return "No file changes are required."

        return _confirmation(
            "write_file",
            file_path,
            "The existing file will be overwritten.",
            mode="overwrite",
            content=content,
            old_preview=_truncate_text(old_content),
            new_preview=_truncate_text(content),
        )

    file_path.parent.mkdir(parents=True, exist_ok=True)
    file_path.write_text(content, encoding="utf-8")

    return f"File created: {file_path}"


def write_file(
    path: str,
    content: str,
    mode: str = "overwrite",
    _source: str = "terminal",
    _user_id: int | None = None,
):
    file_path = validate_write(
        resolve_user_path(path),
        _source,
        _user_id,
    )

    mode = mode.lower().strip()

    if mode not in {"overwrite", "append"}:
        raise ValueError("mode must be overwrite or append.")

    if file_path.suffix.lower() not in TEXT_EXTENSIONS:
        raise ValueError(
            "write_file supports only TXT, MD, PY, LOG, CSV and JSON files."
        )

    if not file_path.exists():
        file_path.parent.mkdir(parents=True, exist_ok=True)
        file_path.write_text(content, encoding="utf-8")
        return f"File created: {file_path}"

    if not file_path.is_file():
        raise ValueError("The provided path is not a file.")

    old_content = file_path.read_text(
        encoding="utf-8",
        errors="replace",
    )

    if mode == "append":
        separator = "\n" if old_content and not old_content.endswith("\n") else ""
        new_content = old_content + separator + content
    else:
        new_content = content

    if new_content == old_content:
        return "No file changes are required."

    return _confirmation(
        "write_file",
        file_path,
        "The file content will be changed.",
        mode=mode,
        content=content,
        old_preview=_truncate_text(old_content),
        new_preview=_truncate_text(new_content),
    )


def read_file(
    path: str,
    _source: str = "terminal",
    _user_id: int | None = None,
) -> str:
    file_path = validate_read(
        resolve_user_path(path),
        _source,
        _user_id,
    )

    if not file_path.exists():
        raise FileNotFoundError(f"File not found: {file_path}")

    if not file_path.is_file():
        raise ValueError("The provided path is not a file.")

    suffix = file_path.suffix.lower()

    if suffix == ".pdf":
        content = _read_pdf(file_path)
    elif suffix == ".docx":
        content = _read_docx(file_path)
    elif suffix == ".xlsx":
        content = _read_xlsx(file_path)
    elif suffix == ".pptx":
        content = _read_pptx(file_path)
    elif suffix == ".json":
        content = _read_json(file_path)
    elif suffix == ".csv":
        content = _read_csv(file_path)
    elif suffix in TEXT_EXTENSIONS:
        content = file_path.read_text(
            encoding="utf-8",
            errors="replace",
        )
    else:
        raise ValueError(f"Unsupported file type: {suffix}")

    if not content.strip():
        return "The file is empty or contains no readable text."

    return content


def search_files(
    query: str = "",
    extension: str = "",
    directory: str = "",
    _source: str = "terminal",
    _user_id: int | None = None,
) -> str:
    roots = (
        [resolve_user_path(directory)]
        if directory
        else _search_roots(_source, _user_id)
    )

    query = query.lower().strip()
    extension = extension.lower().strip()

    if extension and not extension.startswith("."):
        extension = "." + extension

    results = []

    for root in roots:
        try:
            root = validate_read(
                root,
                _source,
                _user_id,
            )
        except PermissionError:
            continue

        if not root.exists():
            continue

        for path in _iter_search_files(root):
            if len(results) >= 30:
                break

            if not path.is_file() or is_sensitive_path(path):
                continue

            try:
                validate_read(
                    path,
                    _source,
                    _user_id,
                )
            except PermissionError:
                continue

            if extension and path.suffix.lower() != extension:
                continue

            if query and query not in path.name.lower():
                continue

            results.append(str(path))

        if len(results) >= 30:
            break

    if not results:
        return "No matching files were found."

    return "\n".join(
        f"{index}. {path}"
        for index, path in enumerate(results, start=1)
    )


def list_directory(
    path: str,
    _source: str = "terminal",
    _user_id: int | None = None,
) -> str:
    directory = validate_read(
        resolve_user_path(path),
        _source,
        _user_id,
    )

    if not directory.exists():
        raise FileNotFoundError(f"Directory not found: {directory}")

    if not directory.is_dir():
        raise ValueError("The provided path is not a directory.")

    output = []

    for item in sorted(
        directory.iterdir(),
        key=lambda item: (not item.is_dir(), item.name.lower()),
    ):
        if is_sensitive_path(item):
            continue

        item_type = "FOLDER" if item.is_dir() else "FILE"
        output.append(f"[{item_type}] {item.name}")

    if not output:
        return "The directory is empty."

    return "\n".join(output)


def open_path(
    path: str,
    _source: str = "terminal",
    _user_id: int | None = None,
) -> str:
    if not _can_use_open_path(_source, _user_id):
        raise PermissionError(
            "Opening files or folders is available only from "
            "the terminal or Web Admin."
        )

    target = validate_read(
        resolve_user_path(path),
        _source,
        _user_id,
    )

    if not target.exists():
        raise FileNotFoundError(f"Path not found: {target}")

    os.startfile(str(target))

    return f"Opened: {target}"


def manage_file(
    action: str,
    path: str,
    destination: str = "",
    _source: str = "terminal",
    _user_id: int | None = None,
):
    source_path = resolve_user_path(path)
    action = action.lower().strip()

    if action == "delete":
        source_path = validate_delete(
            source_path,
            _source,
            _user_id,
        )

        if not source_path.exists():
            raise FileNotFoundError(f"File not found: {source_path}")

        if not source_path.is_file():
            raise ValueError("Delete is supported only for files.")

        return _confirmation(
            "delete",
            source_path,
            "The file will be moved to the Recycle Bin.",
        )

    if action not in {"move", "rename"}:
        raise ValueError("Invalid action. Use delete, move, or rename.")

    source_path = validate_read(
        source_path,
        _source,
        _user_id,
    )
    validate_write(
        source_path,
        _source,
        _user_id,
    )

    if not source_path.exists():
        raise FileNotFoundError(f"File not found: {source_path}")

    if not source_path.is_file():
        raise ValueError("This operation is supported only for files.")

    if not destination:
        raise ValueError("destination is required.")

    if action == "rename":
        new_name = PureWindowsPath(
            destination.strip().strip('"').strip("'")
        ).name

        if not new_name or new_name in {".", ".."}:
            raise ValueError("A valid new file name is required.")

        destination_path = source_path.with_name(new_name)
    else:
        destination_path = resolve_user_path(destination)

        if destination_path.exists() and destination_path.is_dir():
            destination_path /= source_path.name

    destination_path = validate_write(
        destination_path,
        _source,
        _user_id,
    )

    if destination_path == source_path:
        raise ValueError("Source and destination are the same.")

    if destination_path.exists():
        raise FileExistsError(
            "A file with the same name already exists at the destination."
        )

    return _confirmation(
        action,
        source_path,
        (
            "The file will be moved."
            if action == "move"
            else "The file will be renamed."
        ),
        destination=destination_path,
    )


def execute_pending_file_action(
    pending_action: dict,
    _source: str,
    _user_id: int | None,
) -> str:
    action = pending_action.get("action")
    path = resolve_user_path(pending_action["path"])
    expected_fingerprint = pending_action.get("fingerprint")

    if expected_fingerprint:
        if not path.exists():
            raise FileNotFoundError(
                "The file no longer exists while the operation is awaiting confirmation."
            )

        if _file_fingerprint(path) != expected_fingerprint:
            raise RuntimeError(
                "The file changed while awaiting confirmation. "
                "The previous operation was blocked for security."
            )

    if action == "write_file":
        path = validate_write(
            path,
            _source,
            _user_id,
        )

        mode = pending_action.get("mode", "overwrite")
        content = pending_action.get("content", "")

        if mode == "overwrite":
            path.write_text(content, encoding="utf-8")
        elif mode == "append":
            old_content = path.read_text(
                encoding="utf-8",
                errors="replace",
            )
            separator = "\n" if old_content and not old_content.endswith("\n") else ""

            with path.open("a", encoding="utf-8") as file:
                file.write(separator + content)
        else:
            raise ValueError("Invalid write mode.")

        return f"File updated: {path}"

    if action == "delete":
        path = validate_delete(
            path,
            _source,
            _user_id,
        )
        send2trash(str(path))
        return "The file was moved to the Recycle Bin."

    if action in {"move", "rename"}:
        path = validate_read(
            path,
            _source,
            _user_id,
        )
        validate_write(
            path,
            _source,
            _user_id,
        )

        destination = validate_write(
            resolve_user_path(pending_action["destination"]),
            _source,
            _user_id,
        )

        if destination.exists():
            raise FileExistsError(
                "The destination file was created while confirmation was pending."
            )

        destination.parent.mkdir(
            parents=True,
            exist_ok=True,
        )

        if action == "move":
            shutil.move(
                str(path),
                str(destination),
            )
            return f"File moved: {destination}"

        path.rename(destination)
        return f"File renamed: {destination}"

    raise ValueError("Invalid pending operation.")


def send_file(
    path: str,
    _source: str = "terminal",
    _user_id: int | None = None,
) -> str:
    if not _can_send_telegram_file(
        _source,
        _user_id,
    ):
        raise PermissionError(
            "Sending a Telegram file is available only from "
            "the terminal, Telegram Admin, or Web Admin."
        )

    file_path = validate_read(
        resolve_user_path(path),
        _source,
        _user_id,
    )

    if is_sensitive_path(file_path):
        raise PermissionError(
            "This file cannot be sent through Telegram."
        )

    if not file_path.exists():
        raise FileNotFoundError(f"File not found: {file_path}")

    if not file_path.is_file():
        raise ValueError("The path to send is not a file.")

    return f"__SEND_FILE__::{file_path}"


def _read_pdf(path: Path) -> str:
    reader = PdfReader(path)
    output = []

    for page_number, page in enumerate(reader.pages, start=1):
        text = page.extract_text()

        if text:
            output.append(
                f"--- Page {page_number} ---\n{text}"
            )

    return "\n\n".join(output)


def _read_docx(path: Path) -> str:
    document = Document(path)
    output = []

    for paragraph in document.paragraphs:
        text = paragraph.text.strip()

        if text:
            output.append(text)

    for table_number, table in enumerate(document.tables, start=1):
        output.append(f"--- Table {table_number} ---")

        for row in table.rows:
            output.append(
                " | ".join(
                    cell.text.strip()
                    for cell in row.cells
                )
            )

    return "\n".join(output)


def _read_xlsx(path: Path) -> str:
    workbook = load_workbook(
        path,
        data_only=True,
        read_only=True,
    )
    output = []

    try:
        for sheet in workbook.worksheets:
            output.append(f"--- Sheet: {sheet.title} ---")

            for row in sheet.iter_rows(values_only=True):
                output.append(
                    " | ".join(
                        str(value) if value is not None else ""
                        for value in row
                    )
                )
    finally:
        workbook.close()

    return "\n".join(output)


def _read_pptx(path: Path) -> str:
    presentation = Presentation(path)
    output = []

    for slide_number, slide in enumerate(
        presentation.slides,
        start=1,
    ):
        output.append(f"--- Slide {slide_number} ---")

        for shape in slide.shapes:
            if not hasattr(shape, "text"):
                continue

            text = shape.text.strip()

            if text:
                output.append(text)

    return "\n".join(output)


def _read_json(path: Path) -> str:
    with path.open(
        "r",
        encoding="utf-8",
    ) as file:
        data = json.load(file)

    return json.dumps(
        data,
        ensure_ascii=False,
        indent=2,
    )


def _read_csv(path: Path) -> str:
    output = []

    with path.open(
        "r",
        encoding="utf-8-sig",
        errors="replace",
        newline="",
    ) as file:
        reader = csv.reader(file)

        for row in reader:
            output.append(
                " | ".join(row)
            )

    return "\n".join(output)