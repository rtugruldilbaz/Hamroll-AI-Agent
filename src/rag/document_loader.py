import json
from pathlib import Path

from docx import Document

from rag.pdf_loader import load_pdf


SUPPORTED_EXTENSIONS = {
    ".pdf",
    ".docx",
    ".txt",
    ".md",
    ".csv",
    ".json",
}


def load_document(
    path: str,
) -> str:
    file_path = Path(path)

    if not file_path.exists():
        raise FileNotFoundError(
            f"Document not found: {file_path.name}"
        )

    if not file_path.is_file():
        raise ValueError(
            "The provided path is not a file."
        )

    extension = (
        file_path.suffix.lower()
    )

    if extension not in SUPPORTED_EXTENSIONS:
        raise ValueError(
            f"Unsupported document type: {extension}"
        )

    if extension == ".pdf":
        return load_pdf(
            str(file_path)
        ).strip()

    if extension == ".docx":
        document = Document(
            str(file_path)
        )

        parts = [
            paragraph.text.strip()
            for paragraph
            in document.paragraphs
            if paragraph.text.strip()
        ]

        for table in document.tables:
            for row in table.rows:
                cells = [
                    cell.text.strip()
                    for cell in row.cells
                    if cell.text.strip()
                ]

                if cells:
                    parts.append(
                        " | ".join(
                            cells
                        )
                    )

        return "\n".join(
            parts
        ).strip()

    if extension == ".json":
        text = file_path.read_text(
            encoding="utf-8-sig"
        )

        data = json.loads(
            text
        )

        return json.dumps(
            data,
            ensure_ascii=False,
            indent=2,
        )

    return file_path.read_text(
        encoding="utf-8-sig",
        errors="replace",
    ).strip()